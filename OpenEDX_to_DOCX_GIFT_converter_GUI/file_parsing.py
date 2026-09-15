import os
import math

from docx import Document
from docx.shared import Pt
from docx.shared import RGBColor


alphabet = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"


class IncorrectSyntax(Exception):
    pass


def MathJax_border_fix_inv(text, id):
    i = 0
    while i != -1 and i < len(text):
        start_tag = text.find("\\(", i)
        if start_tag != -1:
            end_tag = text.find("\\)", start_tag + 2)
            if end_tag == -1:
                raise IncorrectSyntax("The closing parenthesis is missing.", "Question: " + id)
            # if text.find('\n', start_tag, end_tag) != -1:
            #     raise IncorrectSyntax("Line feed inside formula", "Question: " + id)
            text = text[: start_tag] + "$$" + text[start_tag + 2: end_tag] + "$$" + text[end_tag + 2:]
        i = start_tag
    return text


# НЕ УДАЛЯТь
# def MathJax_border_del_gift(text):
#     i = 0
#     while i != -1 and i < len(text):
#         start_tag = text.find("\\(", i)
#         if start_tag != -1:
#             end_tag = text.find("\\)", start_tag + 2)
#             if end_tag + 2 < len(text):
#                 text = text[: start_tag] + text[start_tag + 2: end_tag] + text[end_tag + 2:]
#             else:
#                 text = text[: start_tag] + text[start_tag + 2: end_tag]
#         i = start_tag
#     return text


def html_replace_inv(text):
    text = text.replace('&amp;', '&')
    text = text.replace('&lt;', '<')
    text = text.replace('&gt;', '>')
    return text


def round_to(num, digits=2):
    if num == 0:
        return 0
    if num - int(num) != 0:
        scale = int(-math.floor(math.log10(abs(num - int(num))))) + digits - 1
        if scale < digits:
            scale = digits
        return round(num, scale)
    else:
        return int(num)


def get_problem_list(lib_xml):
    prob_list = []
    library = open(lib_xml, "r", encoding="utf-8")
    lib_text = library.read()
    library.close()
    next_prob = lib_text.find("<problem url_name=")
    while next_prob != -1:
        opening = lib_text.find("\"", next_prob)
        prob_list.append(lib_text[opening + 1: lib_text.find("\"", opening + 1)])
        next_prob = lib_text.find("<problem url_name=", opening)
    return prob_list


def get_problem_type(text):
    type_found = False
    problem_type = ""
    idx = 0
    types = ['multiplechoiceresponse', 'choiceresponse', 'stringresponse', 'formularesponse', 'numericalresponse']
    for i in range(2):
        for pr_type in types:
            if text.find("<" + pr_type, idx) != -1:
                if not type_found:
                    problem_type = pr_type
                    type_found = True
                    idx = text.find(pr_type) + 1
                else:
                    if i == 1:
                        problem_type = 'mixed'
    if not problem_type:
        if text.find('choicegroup') != -1:
            problem_type = 'multiplechoiceresponse'
        if text.find('checkboxgroup') != -1:
            problem_type = 'choiceresponse'
    print(problem_type)
    return problem_type


# img in the end
# def get_label(text, id):
#     label = ""
#     opening = text.find("<label")
#     extras = []
#     i = 0
#     while text.find("<p>", i) != -1:
#         i = text.find("<p>", i) + 1
#         extras.append(i - 1)
#     image = text.find("<img")
#     if opening != -1:  # label was found
#         label_label = html_replace_inv(MathJax_border_fix_inv(text[text.find(">", opening) + 1: text.find("</label", opening)], id)).strip()
#         if extras:
#             before = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras if extra < opening]
#             after = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras if extra > opening]
#             label = ("\n".join(before) + "\n" + label_label + "\n" + "\n".join(after)).strip(" \n")
#         else:
#             label = label_label
#     else:
#         extra_texts = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras]
#         label = "\n".join(extra_texts)
#     if image != -1:
#         label += "\n" + text[image: text.find("/>", image) + 2]
#     return label


def get_label(text, id, q_type):
    label = ""
    opening = text.find("<label")
    extras = []
    i = 0
    while text.find("<p>", i) != -1 and (text.find("<p>", i) < text.find(q_type) or text.find("<p>", i) <
                                         text.find("choicegroup") or text.find("<p>", i) < text.find("checkboxgroup")):
        i = text.find("<p>", i) + 1
        extras.append(i - 1)
    image = text.find("<img")
    if opening != -1:  # label was found
        label_label = html_replace_inv(MathJax_border_fix_inv(text[text.find(">", opening) + 1: text.find("</label", opening)], id)).strip()
        if extras:
            before = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras if extra < opening]
            after = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras if extra > opening]
            befores = [extra for extra in extras if extra < opening]
            afters = [extra for extra in extras if extra > opening]
            if image != -1:
                if befores and befores[0] < image < befores[len(befores) - 1]:
                    for i in range(len(befores)):
                        if i != len(befores) and befores[i] < image < befores[i + 1]:
                            before.insert(i + 1, text[image: text.find("/>", image) + 2])
                elif afters and afters[0] < image < afters[len(afters) - 1]:
                    for i in range(len(afters)):
                        if i != len(afters) and afters[i] < image < afters[i + 1]:
                            after.insert(i + 1, text[image: text.find("/>", image) + 2])
                elif afters and image < afters[0]:
                    after.insert(0, text[image: text.find("/>", image) + 2])
                elif befores and image > befores[len(befores) - 1]:
                    before.append(text[image: text.find("/>", image) + 2])
                elif afters and afters[len(afters) - 1] < image < opening:
                    after.append(text[image: text.find("/>", image) + 2])
                elif befores:
                    before.insert(0, text[image: text.find("/>", image) + 2])
            label = ("\n".join(before) + "\n" + label_label + "\n" + "\n".join(after)).strip(" \n")
        elif image != -1:
            label = text[image: text.find("/>", image) + 2] + label_label if image < opening else label_label + text[image: text.find("/>", image) + 2]
        else:
            label = label_label
    else:
        extra_texts = [html_replace_inv(MathJax_border_fix_inv(text[extra + 3: text.find("</p>", extra)], id)).strip() for extra in extras]
        label = "\n".join(extra_texts)
    return label


def multiple_choice_question_body(text, id):
    i = 0
    body = "{\n"
    while text.find("<choice ", i) != -1:
        choice_idx = text.find("<choice ", i)
        if text[choice_idx + 17: text.find("\"", choice_idx + 17)] in ("true", "True"):
            body += "="
        else:
            body += "~"
        if text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx)) != -1:
            body += html_replace_inv(MathJax_border_fix_inv(text[text.find(">", choice_idx) + 1: text.find("<choicehint", choice_idx)], id)).strip(" \n")
            hint = text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx))
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[text.find(">", hint) + 1:
                                                                       text.find("</choicehint>", hint)], id)).strip(" \n")
        else:
            body += html_replace_inv(MathJax_border_fix_inv(text[text.find(">", choice_idx) + 1: text.find("</choice>", choice_idx)], id)).strip(" \n")
        body += "\n"
        i = choice_idx + 1
    body += "}"
    return body


def checkbox_question_body(text, id):
    # по умолчанию делим 100% на количество правильных ответов
    body = "{\n"
    correct_num = text.count("<choice correct=\"true\">")  # количество правильных ответов
    percent = int(100 / correct_num)
    add = 0
    if 100 % correct_num != 0:
        add = 100 % correct_num

    i = 0
    ct = 0
    while text.find("<choice ", i) != -1:
        choice_idx = text.find("<choice ", i)
        if text[choice_idx + 17: text.find("\"", choice_idx + 17)] in ("true", "True"):
            body += "~%"
            body += str(percent) + "%" if ct > 0 else str(percent + add) + "%"
            ct += 1
        else:
            body += "~"
        if text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx))  != -1:
            body += html_replace_inv(MathJax_border_fix_inv(text[text.find(">", choice_idx) + 1: text.find("<choicehint", choice_idx)], id)).strip(" \n")
            hint = text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx))
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[text.find(">", hint) + 1:
                                                                       text.find("</choicehint>", hint)], id)).strip(" \n")
        else:
            body += html_replace_inv(MathJax_border_fix_inv(text[text.find(">", choice_idx) + 1: text.find("</choice>", choice_idx)], id)).strip(" \n")
        body += "\n"
        i = choice_idx + 1
    body += "}"
    return body


def text_match_question_body(text, id):
    body = "{\n"
    # main answer
    i = text.find("<stringresponse answer=")
    body += "=" + html_replace_inv(MathJax_border_fix_inv(text[i + 24: text.find("\">", i)], id))
    corr_hint = text.find("<correcthint>")
    if corr_hint != -1:
        body += "#" + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13: text.find("\n", corr_hint)], id))
    body += "\n"
    while text.find("<additional_answer", i) != -1:
        answer_idx = text.find("<additional_answer", i)
        answer = html_replace_inv(MathJax_border_fix_inv(text[answer_idx + 27: text.find("\"/>", answer_idx)], id))
        body += "=" + answer
        hint = text.find("<stringequalhint answer=\"" + answer)
        if hint != -1:
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[text.find(">", hint) + 1:
                                                                        text.find("\n", hint)], id))
        elif corr_hint != -1:
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13: text.find("\n", corr_hint)], id))
        body += "\n"
        i = answer_idx + 1
    body += "}"
    return body


def math_expression_question_body(text, id):
    body = "{\n"
    start = text.find("<formularesponse")
    body += "=" + html_replace_inv(MathJax_border_fix_inv(text[text.find("answer=", start) + 8:
                                                                text.find("\">", start)], id)) + "\n" + "}"
    return body


def numeric_question_body(text, str_tolerance, id):
    body = "{\n"
    start = text.find("<numericalresponse")
    corr_hint = text.find("<correcthint>")
    tolerance = 0
    if int(str_tolerance[:-1]) != 0:
        tolerance = int(str_tolerance[:-1])
    main_answer = html_replace_inv(MathJax_border_fix_inv(text[start + len("numericalresponse answer=\"") + 1:
                                                                text.find("\">", start)], id))
    if text.find("<additional_answer") != -1:
        body += "#\n=" + main_answer
        if tolerance != 0:
            body += ":" + str(abs(round_to(float(main_answer) * (tolerance / 100))))
        if corr_hint != -1:
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13: text.find("\n", corr_hint)], id))
        body += "\n"
        i = 0
        while text.find("<additional_answer", i) != -1:
            answer_idx = text.find("<additional_answer", i)
            answer = html_replace_inv(MathJax_border_fix_inv(text[answer_idx + len("additional_answer answer=\"") + 1:
                                                                   text.find("\"/>", answer_idx)], id))
            body += "=" + answer
            if tolerance != 0:
                body += ":" + str(abs(round_to(float(answer) * (tolerance / 100))))
            if corr_hint != -1:
                body += "#" + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13:
                                                                            text.find("\n", corr_hint)], id))
            body += "\n"
            i = answer_idx + 1
    else:
        body += "#" + main_answer
        if tolerance != 0:
            body += ":" + str(abs(round_to(float(main_answer) * (tolerance / 100))))
        if corr_hint != -1:
            body += "#" + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13: text.find("\n", corr_hint)], id))
        body += "\n"
    body += "}"
    return body


def parse_xml_to_gift(path_to_library, dir_res, f_attr, q_attr):
    if path_to_library != "":
        path_to_policies = ""
        path_to_problem = ""
        lib_xml = ""
        i = 1
        for elem in os.listdir(path_to_library):
            if elem == "policies":
                path_to_policies = path_to_library + "/" + "policies"
            if elem == "problem":
                path_to_problem = path_to_library + "/" + "problem"
            if elem == "library.xml":
                lib_xml = path_to_library + "/" + "library.xml"
        if path_to_policies != "" and path_to_problem != "" and lib_xml != "":
            if len(os.listdir(path_to_problem)) != 0:
                problem_list = get_problem_list(lib_xml)
                gift_content = ""
                try:
                    for file in problem_list:
                        unknown = False
                        xml_problem = open(path_to_problem + "/" + file + ".xml", "r", encoding="utf-8")
                        xml_text = xml_problem.read()
                        xml_problem.close()

                        # question type
                        problem_type = get_problem_type(xml_text)
                        # question name
                        problem_name = q_attr.display_name + " " + str(i) if q_attr.allow_numbering == "true" \
                            else q_attr.display_name
                        i += 1
                        # label
                        problem_label = get_label(xml_text, file[:-4], problem_type)
                        # body
                        problem_body = ""
                        if problem_type == "multiplechoiceresponse":
                            print("found mult")
                            problem_body = multiple_choice_question_body(xml_text, file[:-4])
                        elif problem_type == "choiceresponse":
                            print("found choice")
                            problem_body = checkbox_question_body(xml_text, file[:-4])
                        elif problem_type == "stringresponse":
                            print("found string")
                            problem_body = text_match_question_body(xml_text, file[:-4])
                        elif problem_type == "formularesponse":
                            print("found formula")
                            problem_body = math_expression_question_body(xml_text, file[:-4])
                        elif problem_type == "numericalresponse":
                            print("found num")
                            problem_body = numeric_question_body(xml_text, q_attr.tolerance, file[:-4])
                            # problem_body = ""
                        else:
                            print("Unknown problem type.")
                            unknown = True

                        if not (unknown or problem_type == "mixed"):
                            question = problem_label + "\n" + problem_body + "\n\n"
                            if problem_name != "":
                                question = "::" + problem_name + "::" + question
                        else:
                            question = "\n{}\n\n"
                            if problem_name != "":
                                question = "::" + problem_name + "::" + question
                        gift_content += question
                    if gift_content[-2:] == "\n\n":
                        gift_content = gift_content[: -2]
                    # write to file
                    res_file = open(dir_res + "/" + f_attr.file_name, "w", encoding="utf-8")
                    res_file.write(gift_content)
                    res_file.close()

                    print("Conversion completed")
                except IncorrectSyntax as ex:
                    print("The conversion was stopped for the following reasons:\n", ex.args[0], sep="")
                    if len(ex.args) > 1:
                        for arg in ex.args[1:]:
                            print(arg)
            else:
                print("No questions found!")
        else:
            print("The unsatisfactory structure of the library directory.")
    else:
        print("Empty name of directory.")


# for both choice types
def choice_question_body_docx(text, id):
    i = 0
    idx = 0  # индексы массива
    ct = 0  # подсчёт для алфавита
    body = []
    while text.find("<choice ", i) != -1:
        choice_idx = text.find("<choice ", i)
        if text[choice_idx + 17: text.find("\"", choice_idx + 17)] in ("true", "True"):
            body.append("*" + alphabet[ct] + ": ")
        else:
            body.append(alphabet[ct] + ": ")

        if text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx)) != -1:
            choice_text = text[text.find(">", choice_idx) + 1: text.find("<choicehint", choice_idx)]
            if choice_text.find("<p>") != -1:
                choice_text = choice_text[choice_text.find("<p>") + 3: choice_text.find("</p>")]
            body[idx] += html_replace_inv(MathJax_border_fix_inv(choice_text, id)).strip()
            idx += 1
            ct += 1
            hint = text.find("<choicehint", choice_idx, text.find("</choice>", choice_idx))
            body.append("Feedback: " + html_replace_inv(MathJax_border_fix_inv(text[text.find(">", hint) + 1:
                                                                                    text.find("</choicehint>", hint)], id)).strip())
            idx += 1
        else:
            choice_text = text[text.find(">", choice_idx) + 1: text.find("</choice>", choice_idx)]
            if choice_text.find("<p>") != -1:
                choice_text = choice_text[choice_text.find("<p>") + 3: choice_text.find("</p>")]
            body[idx] += html_replace_inv(MathJax_border_fix_inv(choice_text, id)).strip()
            idx += 1
            ct += 1
        i = choice_idx + 1
    return body


def text_match_question_body_docx(text, id):
    body = []
    ct = 1
    # main answer
    i = text.find("<stringresponse answer=")
    body.append("*A: " + html_replace_inv(MathJax_border_fix_inv(text[i + 24: text.find("\">", i)], id)))
    corr_hint = text.find("<correcthint>")
    if corr_hint != -1:
        body.append("Feedback: " + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13:
                                                                                text.find("\n", corr_hint)], id)))
    while text.find("<additional_answer", i) != -1:
        answer_idx = text.find("<additional_answer", i)
        answer = html_replace_inv(MathJax_border_fix_inv(text[answer_idx + 27: text.find("\"/>", answer_idx)], id))
        body.append("*" + alphabet[ct] + ": " + answer)
        ct += 1
        i = answer_idx + 1
    i = 0
    # wrong answers
    while text.find("<stringequalhint", i) != -1:
        hint_idx = text.find("<stringequalhint", i)
        wrong_answer = html_replace_inv(MathJax_border_fix_inv(text[hint_idx + len("<stringequalhint answer=\""):
                                                                    text.find("\">", hint_idx)], id))
        body.append(alphabet[ct] + ": " + wrong_answer)
        body.append("Feedback: " + html_replace_inv(MathJax_border_fix_inv(text[text.find("\">", hint_idx) + 2:
                                                                                text.find("\n", hint_idx)], id)))
        ct += 1
        i = hint_idx + 1
    return body


def math_expression_question_body_docx(text, id):
    body = []
    start = text.find("<formularesponse")
    body.append("*A: " + html_replace_inv(MathJax_border_fix_inv(text[text.find("answer=", start) + 8:
                                                             text.find("\">", start)], id)))
    return body


def numeric_question_body_docx(text, id):
    body = []
    start = text.find("<numericalresponse")
    corr_hint = text.find("<correcthint>")
    main_answer = html_replace_inv(MathJax_border_fix_inv(text[start + len("numericalresponse answer=\"") + 1:
                                                               text.find("\">", start)], id))
    body.append("*A: " + main_answer)
    if corr_hint != -1:
        body.append("Feedback: " + html_replace_inv(MathJax_border_fix_inv(text[corr_hint + 13:
                                                                                text.find("\n", corr_hint)], id)))
    if text.find("<additional_answer") != -1:
        i = 0
        ct = 1
        while text.find("<additional_answer", i) != -1:
            answer_idx = text.find("<additional_answer", i)
            answer = html_replace_inv(MathJax_border_fix_inv(text[answer_idx + len("additional_answer answer=\"") + 1:
                                                                  text.find("\"/>", answer_idx)], id))
            body.append("*" + alphabet[ct] + ": " + answer)
            ct += 1
            i = answer_idx + 1
    return body


def parse_xml_to_docx(path_to_library, dir_res, f_attr):
    if path_to_library != "":
        path_to_policies = ""
        path_to_problem = ""
        lib_xml = ""
        i = 1
        for elem in os.listdir(path_to_library):
            if elem == "policies":
                path_to_policies = path_to_library + "/" + "policies"
            if elem == "problem":
                path_to_problem = path_to_library + "/" + "problem"
            if elem == "library.xml":
                lib_xml = path_to_library + "/" + "library.xml"
        if path_to_policies != "" and path_to_problem != "" and lib_xml != "":
            if len(os.listdir(path_to_problem)) != 0:
                problem_list = get_problem_list(lib_xml)
                try:
                    doc = Document()
                    header = doc.add_heading('Practice quiz', level=1)
                    header.runs[0].font.bold = True
                    header.runs[0].font.size = Pt(12)
                    header.runs[0].font.name = 'Open Sans'
                    header.runs[0].font.color.rgb = RGBColor(0, 0, 0)
                    paragraph = doc.add_paragraph().add_run(f_attr.test_name)  # название теста
                    font = paragraph.font
                    font.size = Pt(11)
                    font.name = 'Arial'
                    font.bold = False
                    doc.add_paragraph()
                    for file in problem_list:
                        unknown = False
                        xml_problem = open(path_to_problem + "/" + file + ".xml", "r", encoding="utf-8")
                        xml_text = xml_problem.read()
                        xml_problem.close()
                        print(file)

                        # question type
                        problem_type = get_problem_type(xml_text)
                        # question name
                        problem_name = "Question " + str(i) + " - "
                        i += 1
                        # label
                        problem_label = get_label(xml_text, file[:-4], problem_type)
                        # body
                        problem_body = []
                        if problem_type == "multiplechoiceresponse":
                            problem_name += "single correct answer"
                            if xml_text.find("shuffle=\"true\"") != -1:
                                problem_name += ", shuffle"
                            # print("found mult")
                            problem_body = choice_question_body_docx(xml_text, file[:-4])
                        elif problem_type == "choiceresponse":
                            problem_name += "checkbox"
                            if xml_text.find("partial_credit=\"EDC\"") != -1:
                                problem_name += ", partial credit"
                            # print("found choice")
                            problem_body = choice_question_body_docx(xml_text, file[:-4])
                        elif problem_type == "stringresponse":
                            problem_name += "text match"
                            # print("found string")
                            problem_body = text_match_question_body_docx(xml_text, file[:-4])
                        elif problem_type == "formularesponse":
                            problem_name += "math expression"
                            # print("found formula")
                            problem_body = math_expression_question_body_docx(xml_text, file[:-4])
                        elif problem_type == "numericalresponse":
                            problem_name += "numeric"
                            # print("found num")
                            problem_body = numeric_question_body_docx(xml_text, file[:-4])
                        elif problem_type == "mixed":
                            problem_name += "mixed"
                        else:
                            print("Unknown problem type.")
                            unknown = True

                        trailing_text = ""
                        if problem_type in ["stringresponse", "formularesponse", "numericalresponse"]:
                            if xml_text.find("<textline") != -1:
                                tr_text_idx = xml_text.find("trailing_text=", xml_text.find("<textline"))
                                trailing_text = html_replace_inv(MathJax_border_fix_inv(
                                    xml_text[tr_text_idx + len("trailing_text=\""): xml_text.find("\"/>", tr_text_idx)],
                                    file[:-4]))

                        if not (unknown or problem_type == "mixed"):
                            # write to doc (one question)
                            paragraph = doc.add_paragraph().add_run(problem_name)
                            font = paragraph.font
                            font.size = Pt(11)
                            font.name = 'Open Sans'
                            font.bold = True
                            paragraph = doc.add_paragraph().add_run(problem_label)
                            font = paragraph.font
                            font.size = Pt(11)
                            font.name = 'Open Sans'
                            font.bold = False
                            for string in problem_body:
                                if string.find("Feedback") == -1:
                                    doc.add_paragraph()
                                paragraph = doc.add_paragraph().add_run(string)
                                font = paragraph.font
                                font.size = Pt(11)
                                font.name = 'Open Sans'
                                font.bold = False
                            if trailing_text:
                                paragraph = doc.add_paragraph().add_run("#" + trailing_text)
                                font = paragraph.font
                                font.size = Pt(11)
                                font.name = 'Open Sans'
                                font.bold = False
                            doc.add_paragraph()
                        else:
                            if problem_type != "mixed":
                                paragraph = doc.add_paragraph().add_run("Question " + str(i - 1))
                            else:
                                paragraph = doc.add_paragraph().add_run(problem_name)
                            font = paragraph.font
                            font.size = Pt(11)
                            font.name = 'Open Sans'
                            font.bold = True
                            paragraph = doc.add_paragraph().add_run(xml_text)
                            font = paragraph.font
                            font.size = Pt(11)
                            font.name = 'Open Sans'
                            font.bold = False
                    doc.save(dir_res + "/" + f_attr.file_name)
                    print("Conversion completed")
                except IncorrectSyntax as ex:
                    print("The conversion was stopped for the following reasons:\n", ex.args[0], sep="")
                    if len(ex.args) > 1:
                        for arg in ex.args[1:]:
                            print(arg)
            else:
                print("No questions found!")
        else:
            print("The unsatisfactory structure of the library directory.")
    else:
        print("The directory does not contain a library.")